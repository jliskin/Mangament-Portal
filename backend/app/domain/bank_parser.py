"""Bank statement parser (CSV / XLSX) with automatic format detection.

Real-world statements differ per bank in: encoding, delimiter, preamble rows
before the header, header language, a single signed amount column vs.
separate credit/debit columns, decimal convention and date order. Rather than
hard-coding one bank, the parser *detects* each of these from the file and
reports what it detected (``ParseResult.report``) so the user can verify it.

Row-level failures never abort the import: they are collected in
``ParseResult.errors`` and shown to the user.
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from .money import AmountParseError, detect_decimal_separator, parse_amount
from .text import normalize, strip_accents

# --------------------------------------------------------------------------
# Column synonyms (es / pt / en / he). Matching is accent- and case-insensitive.
# Order matters: the first field claiming a column wins.
# --------------------------------------------------------------------------
COLUMN_SYNONYMS: dict[str, list[str]] = {
    "value_date": ["fecha valor", "f valor", "value date", "data valor", "תאריך ערך"],
    "txn_date": ["fecha", "fecha operacion", "fecha movimiento", "fecha contable", "date",
                 "transaction date", "posting date", "booking date", "data", "data lancamento",
                 "data movimento", "תאריך"],
    "credit": ["credito", "creditos", "abono", "abonos", "haber", "ingreso", "ingresos", "credit",
               "credits", "deposit", "deposits", "money in", "paid in", "entrada", "entradas",
               "זכות"],
    "debit": ["debito", "debitos", "cargo", "cargos", "debe", "egreso", "egresos", "debit",
              "debits", "withdrawal", "withdrawals", "money out", "paid out", "saida", "saidas",
              "חובה"],
    "amount": ["importe", "monto", "amount", "valor", "valor r", "cantidad", "סכום"],
    "direction": ["tipo", "type", "d c", "dc", "cr dr", "naturaleza", "natureza", "sentido"],
    "balance": ["saldo", "balance", "running balance", "saldo disponible", "יתרה"],
    "reference": ["referencia", "ref", "reference", "documento", "nro documento", "n documento",
                  "comprobante", "numero", "documento numero", "check number", "אסמכתא"],
    "counterparty": ["ordenante", "remitente", "nombre", "beneficiario", "contraparte", "titular",
                     "payer", "sender", "name", "counterparty", "pagador", "favorecido"],
    "description": ["descripcion", "concepto", "detalle", "detalles", "movimiento", "glosa",
                    "description", "memo", "details", "narrative", "historico", "lancamento",
                    "תיאור", "פרטים", "תיאור התנועה"],
}

DATE_FORMATS_DAYFIRST = ["%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%d-%m-%y", "%d.%m.%y"]
DATE_FORMATS_MONTHFIRST = ["%m/%d/%Y", "%m-%d-%Y", "%m/%d/%y"]
DATE_FORMATS_ISO = ["%Y-%m-%d", "%Y/%m/%d", "%Y%m%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S"]

# Patterns that isolate the payer's name inside free-text descriptions.
COUNTERPARTY_PATTERNS = [
    re.compile(r"(?:transf(?:erencia)?\.?\s*(?:recibida|inmediata)?\s*(?:de|desde)?|"
               r"pix\s+recebido(?:\s+de)?|ted\s+recebida(?:\s+de)?|"
               r"(?:transfer|payment)\s+from|dep(?:osito)?\.?\s+de|de:|from:)\s+"
               r"(?P<name>[a-z][a-z .'\-]{3,60}?)(?:\s+(?:ref|cbu|cpf|cuit|rut|nif|iban|concepto|cta|apto|apt|depto|dpto|unidad|unit|piso|local|alquiler|renta|aluguel|rent|expensas|cuota)\b|[\d/#:|]|$)",
               re.IGNORECASE),
]


class StatementParseError(ValueError):
    """The file as a whole cannot be interpreted (no header, no amount column...)."""


@dataclass
class ParsedTransaction:
    row_number: int
    txn_date: date
    amount_minor: int
    description: str
    counterparty_name: str | None = None
    reference: str | None = None
    value_date: date | None = None
    balance_minor: int | None = None
    raw: dict[str, str] = field(default_factory=dict)
    fingerprint: str = ""

    @property
    def is_credit(self) -> bool:
        return self.amount_minor > 0


@dataclass
class RowError:
    row_number: int
    message: str
    raw: dict[str, str]


@dataclass
class ParseResult:
    transactions: list[ParsedTransaction]
    errors: list[RowError]
    report: dict

    @property
    def credits(self) -> list[ParsedTransaction]:
        return [t for t in self.transactions if t.is_credit]


# --------------------------------------------------------------------------
# Raw table readers
# --------------------------------------------------------------------------
ENCODINGS = ("utf-8-sig", "cp1252", "cp1255", "latin-1")


def _read_csv(content: bytes) -> tuple[list[list[str]], dict]:
    text, used = None, None
    for enc in ENCODINGS:
        try:
            text, used = content.decode(enc), enc
            break
        except UnicodeDecodeError:
            continue
    if text is None:  # latin-1 never fails, but keep the type checker honest
        raise StatementParseError("No se pudo decodificar el archivo")
    sample = "\n".join(text.splitlines()[:40])
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = max(";,\t|", key=sample.count)
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    return rows, {"encoding": used, "delimiter": delimiter}


def _read_xlsx(content: bytes) -> tuple[list[list[object]], dict]:
    from openpyxl import load_workbook  # local import: optional dependency path

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = [["" if v is None else v for v in row] for row in ws.iter_rows(values_only=True)]
    wb.close()
    return rows, {"sheet": ws.title}


def read_table(content: bytes, filename: str) -> tuple[list[list], dict]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")) or content[:2] == b"PK":
        return _read_xlsx(content)
    if name.endswith(".xls"):
        raise StatementParseError("Formato .xls antiguo: guárdelo como .xlsx o .csv desde Excel")
    return _read_csv(content)


# --------------------------------------------------------------------------
# Header & column detection
# --------------------------------------------------------------------------
def _norm_header(cell) -> str:
    s = strip_accents(str(cell or "")).lower()
    s = re.sub(r"[^a-z0-9א-ת ]+", " ", s)
    return " ".join(s.split())


def _field_for_header(h: str) -> str | None:
    if not h:
        return None
    # exact synonym first, then "starts with" (e.g. 'importe eur', 'saldo $')
    for fld, syns in COLUMN_SYNONYMS.items():
        if h in syns:
            return fld
    for fld, syns in COLUMN_SYNONYMS.items():
        if any(h.startswith(s + " ") for s in syns):
            return fld
    return None


def detect_header_row(rows: list[list], scan: int = 30) -> int:
    best_idx, best_hits = -1, 0
    for i, row in enumerate(rows[:scan]):
        hits = len({f for f in (_field_for_header(_norm_header(c)) for c in row) if f})
        if hits > best_hits:
            best_idx, best_hits = i, hits
    if best_hits < 2:
        raise StatementParseError(
            "No se encontró la fila de encabezados (se esperaban columnas como Fecha / Importe / Concepto)")
    return best_idx


def map_columns(header: list) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(header):
        fld = _field_for_header(_norm_header(cell))
        if fld and fld not in mapping:
            mapping[fld] = idx
    if "txn_date" not in mapping and "value_date" in mapping:
        mapping["txn_date"] = mapping.pop("value_date")
    if "txn_date" not in mapping:
        raise StatementParseError("No se encontró columna de fecha")
    if "amount" not in mapping and not ({"credit", "debit"} & mapping.keys()):
        raise StatementParseError("No se encontró columna de importe (ni crédito/débito)")
    return mapping


# --------------------------------------------------------------------------
# Value parsing
# --------------------------------------------------------------------------
EXCEL_EPOCH = date(1899, 12, 30)


def detect_date_order(samples: list) -> str:
    """'dayfirst' | 'monthfirst' | 'iso' — decided over the whole column."""
    strings = [str(s).strip() for s in samples if isinstance(s, str) and str(s).strip()]
    if strings and all(re.match(r"^\d{4}[-/]?\d{2}", s) for s in strings):
        return "iso"
    first_gt12 = second_gt12 = False
    for s in strings:
        m = re.match(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.]\d{2,4}", s)
        if m:
            a, b = int(m.group(1)), int(m.group(2))
            first_gt12 |= a > 12
            second_gt12 |= b > 12
    if second_gt12 and not first_gt12:
        return "monthfirst"
    return "dayfirst"  # default for non-US banks; unambiguous evidence overrides it


def parse_date(raw, order: str = "dayfirst") -> date:
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    if isinstance(raw, (int, float)) and 20000 < raw < 80000:  # Excel serial
        return EXCEL_EPOCH + timedelta(days=int(raw))
    s = str(raw or "").strip()
    if not s:
        raise ValueError("fecha vacía")
    s = s.split(" ")[0] if re.match(r"^\d{1,2}[/\-.]", s) else s
    formats = {
        "dayfirst": DATE_FORMATS_DAYFIRST + DATE_FORMATS_ISO,
        "monthfirst": DATE_FORMATS_MONTHFIRST + DATE_FORMATS_ISO,
        "iso": DATE_FORMATS_ISO + DATE_FORMATS_DAYFIRST,
    }[order]
    for fmt_ in formats:
        try:
            return datetime.strptime(s, fmt_).date()
        except ValueError:
            continue
    raise ValueError(f"fecha no reconocida: {raw!r}")


def extract_counterparty(description: str) -> str | None:
    for pat in COUNTERPARTY_PATTERNS:
        m = pat.search(strip_accents(description or ""))
        if m:
            name = " ".join(m.group("name").split()).strip(" .-")
            if len(name) >= 4:
                return name.title()
    return None


def _cell(row: list, idx: int | None):
    if idx is None or idx >= len(row):
        return ""
    v = row[idx]
    return v.strip() if isinstance(v, str) else v


def _is_blank(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


DEBIT_MARKERS = {"d", "db", "dr", "debito", "debit", "cargo", "egreso", "saida", "-"}


# --------------------------------------------------------------------------
# Main entry point
# --------------------------------------------------------------------------
def parse_statement(content: bytes, filename: str, account_key: str = "") -> ParseResult:
    rows, meta = read_table(content, filename)
    if not any(any(not _is_blank(c) for c in r) for r in rows):
        raise StatementParseError("El archivo está vacío")
    h = detect_header_row(rows)
    header = rows[h]
    cols = map_columns(header)
    body = rows[h + 1:]

    def column_values(fld: str) -> list:
        return [_cell(r, cols.get(fld)) for r in body if cols.get(fld) is not None]

    amount_samples = [str(v) for f in ("amount", "credit", "debit", "balance")
                      for v in column_values(f) if isinstance(v, str) and v]
    decimal_sep = detect_decimal_separator(amount_samples)
    date_order = detect_date_order(column_values("txn_date"))

    txns: list[ParsedTransaction] = []
    errors: list[RowError] = []
    skipped = 0
    for offset, row in enumerate(body):
        row_no = h + 2 + offset  # 1-based, as the user sees it in Excel
        if not any(not _is_blank(c) for c in row):
            continue
        raw = {str(header[i]): str(v) for i, v in enumerate(row) if i < len(header) and not _is_blank(v)}
        date_cell = _cell(row, cols["txn_date"])
        amount_cells = [_cell(row, cols.get(f)) for f in ("amount", "credit", "debit") if f in cols]
        if _is_blank(date_cell) or (not re.search(r"\d", str(date_cell))
                                    and all(_is_blank(v) for v in amount_cells)):
            skipped += 1  # footer / subtotal lines such as "Saldo final"
            continue
        try:
            txn_date = parse_date(date_cell, date_order)
            amount = _row_amount(row, cols, decimal_sep)
            if amount == 0:
                continue
            description = str(_cell(row, cols.get("description")) or "")
            counterparty = str(_cell(row, cols.get("counterparty")) or "") or extract_counterparty(description)
            reference = str(_cell(row, cols.get("reference")) or "") or None
            value_date = None
            if cols.get("value_date") is not None and not _is_blank(_cell(row, cols["value_date"])):
                value_date = parse_date(_cell(row, cols["value_date"]), date_order)
            balance = None
            if cols.get("balance") is not None and not _is_blank(_cell(row, cols["balance"])):
                balance = parse_amount(_cell(row, cols["balance"]), decimal_sep)
        except (ValueError, AmountParseError) as exc:
            errors.append(RowError(row_no, str(exc), raw))
            continue
        txns.append(ParsedTransaction(
            row_number=row_no, txn_date=txn_date, value_date=value_date, amount_minor=amount,
            description=description, counterparty_name=counterparty or None, reference=reference,
            balance_minor=balance, raw=raw,
        ))

    assign_fingerprints(txns, account_key)
    report = {
        **meta,
        "header_row": h + 1,
        "columns": {fld: str(header[i]) for fld, i in cols.items()},
        "decimal_separator": decimal_sep,
        "date_order": date_order,
        "rows_parsed": len(txns),
        "rows_error": len(errors),
        "rows_skipped": skipped,
        "credits": sum(1 for t in txns if t.is_credit),
        "debits": sum(1 for t in txns if not t.is_credit),
    }
    return ParseResult(txns, errors, report)


def _row_amount(row: list, cols: dict[str, int], decimal_sep: str) -> int:
    if "credit" in cols or "debit" in cols:
        credit_raw, debit_raw = _cell(row, cols.get("credit")), _cell(row, cols.get("debit"))
        credit = abs(parse_amount(credit_raw, decimal_sep)) if not _is_blank(credit_raw) else 0
        debit = abs(parse_amount(debit_raw, decimal_sep)) if not _is_blank(debit_raw) else 0
        if credit or debit or "amount" not in cols:
            return credit - debit
    amount = parse_amount(_cell(row, cols["amount"]), decimal_sep)
    if "direction" in cols:
        marker = _norm_header(_cell(row, cols["direction"]))
        if marker in DEBIT_MARKERS or marker.startswith("deb"):
            amount = -abs(amount)
        elif marker:
            amount = abs(amount)
    return amount


def assign_fingerprints(txns: list[ParsedTransaction], account_key: str) -> None:
    """Deterministic identity of a bank row, stable across re-imports.

    Two genuinely identical rows on the same day (same amount, same text) are
    disambiguated by their occurrence index, so they are both kept, while the
    same statement imported twice produces identical fingerprints (dedup).
    """
    seen: dict[str, int] = {}
    for t in txns:
        base = "|".join([
            account_key, t.txn_date.isoformat(), str(t.amount_minor),
            normalize(t.description), (t.reference or "").strip().lower(),
            "" if t.balance_minor is None else str(t.balance_minor),
        ])
        n = seen.get(base, 0)
        seen[base] = n + 1
        t.fingerprint = hashlib.sha256(f"{base}|{n}".encode()).hexdigest()
